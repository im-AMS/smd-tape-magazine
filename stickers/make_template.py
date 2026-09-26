"""Build parts-template.xlsx: the parts list as a spreadsheet, with dropdowns,
header notes and checks, prefilled with example.csv.

    uv run make_template.py            # writes parts-template.xlsx next to this file

labels.py reads the "parts" sheet of the workbook directly (or a CSV).
"""
import csv, pathlib
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from labels import COLUMNS, KIND_TO_CLASS, SUPPLIERS, XLSX_SHEET

HERE = pathlib.Path(__file__).parent
ROWS = 1000                     # rows that get dropdowns and checks

NOTES = {
    "kind": "Required. Picks the sticker colour: R, C, L, LED/D/ESD/TVS, Q/FET, IC/U, J/CONN. "
            "Anything else prints in grey (misc).",
    "value": "Required. 330, 10k, 100n, 10u, a colour for LEDs, or the part number for ICs. "
             "Never shortened: a value that doesn't fit its sticker stops the run with an error.",
    "tolerance": "Optional. e.g. 1%, 10%.",
    "rating": "Optional. Resistors: power, e.g. 100m (= 100 mW). Capacitors: voltage and "
              "dielectric, e.g. 50v X7R.",
    "subtype": "Optional. e.g. MLCC.",
    "footprint": "Optional. Package printed in the band on top of the front sticker, e.g. 0603.",
    "supplier": "Optional. LCSC, DigiKey or Mouser: the QR link is built from the sku. "
                "Any other shop: pick Other (or type its name) and put the link in url. "
                "Empty = LCSC.",
    "sku": "Optional. The supplier's part number, e.g. C23138 (LCSC), 497-5235-1-ND "
           "(DigiKey), 511-USBLC6-2SC6 (Mouser). Printed on the strip.",
    "url": "Optional. Product link, e.g. robu.in/?s=R178998&post_type=product (https:// not "
           "needed: shorter is better for the QR). Needed for suppliers other than "
           "LCSC/DigiKey/Mouser; if filled for those, it replaces the built-in link. Long links "
           "make a denser QR: labels.py warns when it may not scan.",
    "quantity": "Optional, for your own stock keeping. NOT printed on the sticker.",
    "feeder": "Tape width in mm: 8, 12, 16, 24 … any number works. Empty = the --feeder "
              "default (8).",
}
WIDTHS = {"kind": 8, "value": 26, "tolerance": 11, "rating": 14, "subtype": 10,
          "footprint": 12, "supplier": 11, "sku": 17, "url": 40, "quantity": 10, "feeder": 9}
LISTS = {
    "kind": list(dict.fromkeys(KIND_TO_CLASS)),
    "feeder": ["8", "12", "16", "24", "32", "44"],
    "supplier": [name for name, _ in SUPPLIERS.values()] + ["Other"],
    "tolerance": ["0.1%", "0.5%", "1%", "2%", "5%", "10%", "20%"],
    "footprint": ["0201", "0402", "0603", "0805", "1206", "1210", "2512", "SOD-123", "SOD-323",
                  "SOT-23", "SOT-23-5", "SOT-23-6", "SOT-223", "SOIC-8", "TSSOP-8", "QFN"],
}
WARN = {
    "kind": "Not a known kind. That's allowed: it prints as misc (grey).",
    "feeder": "Not one of the usual widths. That's allowed: enter the tape width in mm.",
    "tolerance": "Not in the list. That's allowed if it's what you want printed.",
    "footprint": "Not in the list. That's allowed: it's printed as you typed it.",
    "supplier": "Not a supplier with a built-in link. That's allowed: put the product link in "
                "the url column.",
}


def col(name):
    return "ABCDEFGHIJKLMNOPQRSTUVWXYZ"[COLUMNS.index(name)]


def build(out: pathlib.Path):
    wb = Workbook()
    ws = wb.active
    ws.title = XLSX_SHEET
    head_fill = PatternFill("solid", fgColor="2B3540")
    for i, name in enumerate(COLUMNS, 1):
        c = ws.cell(row=1, column=i, value=name)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = head_fill
        c.alignment = Alignment(horizontal="center")
        c.comment = Comment(NOTES[name], "parts template", width=260, height=110)
        ws.column_dimensions[col(name)].width = WIDTHS[name]
    ws.freeze_panes = "A2"

    # Every data cell is text, so Excel keeps 1% as "1%" and 0603 as "0603".
    for r in range(2, ROWS + 2):
        for i in range(1, len(COLUMNS) + 1):
            ws.cell(row=r, column=i).number_format = "@"

    with open(HERE / "example.csv", newline="") as fh:
        for r, row in enumerate(csv.DictReader(fh), 2):
            for name in COLUMNS:
                ws.cell(row=r, column=COLUMNS.index(name) + 1, value=row.get(name, ""))

    for name, items in LISTS.items():
        dv = DataValidation(type="list", formula1='"' + ",".join(items) + '"', allow_blank=True,
                            showErrorMessage=True, errorStyle="warning",
                            errorTitle="Check this value", error=WARN[name],
                            showInputMessage=False)
        dv.add(f"{col(name)}2:{col(name)}{ROWS + 1}")
        ws.add_data_validation(dv)

    # a supplier without a built-in link needs a url (when there is a sku to link)
    # (Excel rejects array constants like {"a","b"} in conditional formatting and
    #  silently drops every rule on the sheet, so spell the comparison out.)
    sup, sku, url = col("supplier"), col("sku"), col("url")
    names = [n for n, _ in SUPPLIERS.values()] + ["Digi-Key"]
    known = ",".join(f'{sup}2="{n}"' for n in names)       # = is case-insensitive in Excel
    ws.conditional_formatting.add(
        f"{url}2:{url}{ROWS + 1}",
        FormulaRule(formula=[f'AND(LEN({sup}2)>0,NOT(OR({known})),LEN({url}2)=0,LEN({sku}2)>0)'],
                    fill=PatternFill("solid", fgColor="F4C7C3")))

    # kind and value are required together: highlight a row that has one without the other
    red = PatternFill("solid", fgColor="F4C7C3")
    k, v = col("kind"), col("value")
    for target, other in ((k, v), (v, k)):
        ws.conditional_formatting.add(
            f"{target}2:{target}{ROWS + 1}",
            FormulaRule(formula=[f'AND(LEN({target}2)=0,LEN({other}2)>0)'], fill=red))

    how = wb.create_sheet("how to use")
    how.column_dimensions["A"].width = 100
    lines = [
        ("SMD magazine stickers: parts list", True),
        ("", False),
        ("1. Fill in the 'parts' sheet, one row per part. The example rows can be deleted.", False),
        ("2. Only kind and value are required. Hover over a column header for what it means.", False),
        ("3. Dropdowns suggest common values; anything else is allowed (you get a warning).", False),
        ("4. Red cells: a row with a kind but no value (or the other way round), or a", False),
        ("   supplier without a built-in link (not LCSC/DigiKey/Mouser) and no url.", False),
        ("   quantity is for your own records: it is never printed.", False),
        ("5. Save, then run:  uv run labels.py parts-template.xlsx --out out --pdf", False),
        ("", False),
        ("Keep the sheet name 'parts' and the header row. Columns can be reordered; extra", False),
        ("columns are ignored. A CSV with the same header row works too.", False),
    ]
    for r, (t, bold) in enumerate(lines, 1):
        how.cell(row=r, column=1, value=t).font = Font(bold=bold, size=13 if bold else 11)
    wb.save(out)


if __name__ == "__main__":
    out = HERE / "parts-template.xlsx"
    build(out)
    print(out)
